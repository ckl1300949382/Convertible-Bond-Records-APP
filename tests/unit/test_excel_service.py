"""Excel 导入导出服务测试（§4 / XLS-01～10）。openpyxl 内存构造工作簿。"""

from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from bond_profit.application.excel_service import ExcelError, ExcelService
from bond_profit.application.sales_service import SalesService
from bond_profit.domain.models import SaleRecord
from bond_profit.infrastructure.database import connect
from bond_profit.infrastructure.repository import RepositoryError, SalesRepository

_HEADERS = ["转债名称", "转债代码", "卖出日期", "卖出价格", "卖出数量", "备注"]


@pytest.fixture()
def service(tmp_path):
    conn = connect(tmp_path / "app.db")
    yield SalesService(SalesRepository(conn))
    conn.close()


@pytest.fixture()
def excel(service):
    return ExcelService(service)


def save_rows(tmp_path, rows, name="import.xlsx", headers=None, sheet_title="卖出记录"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append(headers if headers is not None else _HEADERS)
    for row in rows:
        sheet.append(row)
    path = tmp_path / name
    workbook.save(path)
    return str(path)


def valid_row(**overrides):
    values = ["平银转债", "012345", date(2026, 1, 15), Decimal("130.00"), 10, ""]
    for key, value in overrides.items():
        values[_HEADERS.index(key)] = value
    return values


# ---- 模板（§4.1） ------------------------------------------------------------

def test_template_generation(tmp_path, excel):
    path = tmp_path / "template.xlsx"
    excel.generate_template(str(path))

    workbook = load_workbook(path)
    assert "卖出记录" in workbook.sheetnames
    assert "填写说明" in workbook.sheetnames
    sheet = workbook["卖出记录"]
    assert [cell.value for cell in sheet[1]] == [
        "转债名称", "转债代码", "卖出日期", "卖出价格", "卖出数量", "备注"
    ]
    assert sheet.cell(row=2, column=2).number_format == "@"  # 代码列文本格式
    assert all(
        cell.value is None for row in sheet.iter_rows(min_row=2) for cell in row
    )  # 数据区不放示例行
    guide = workbook["填写说明"]
    assert guide.cell(row=1, column=1).value == "填写说明"


# ---- 解析与预览（§4.2 / §4.4） -----------------------------------------------

def test_preview_valid_rows_both_date_forms(tmp_path, excel):
    """XLS-07：Excel 日期单元格与 ISO 文本日期均可识别。"""
    path = save_rows(
        tmp_path,
        [valid_row(**{"卖出日期": date(2026, 1, 15)}), valid_row(**{"卖出日期": "2026-02-01"})],
    )
    preview = excel.preview_import(path)
    assert preview.error_count == 0
    assert preview.new_count == 2
    assert preview.rows[0].record.sale_date == date(2026, 1, 15)
    assert preview.rows[1].record.sale_date == date(2026, 2, 1)


def test_xls01_in_file_duplicates_skip(tmp_path, excel):
    path = save_rows(tmp_path, [valid_row(), valid_row()])
    preview = excel.preview_import(path)
    assert preview.new_count == 1
    assert preview.duplicate_count == 1
    assert preview.rows[0].status == "new"
    assert preview.rows[1].status == "duplicate"


def test_xls02_existing_db_duplicate_skips(tmp_path, excel, service):
    service.create_record(SaleRecord("平银转债", "012345", date(2026, 1, 15), Decimal("130.00"), 10))
    path = save_rows(tmp_path, [valid_row()])
    preview = excel.preview_import(path)
    assert preview.new_count == 0
    assert preview.duplicate_count == 1


def test_xls03_note_blank_variants_same_key(tmp_path, excel):
    """XLS-03：空备注、空字符串、纯空格去重视为相同。"""
    path = save_rows(
        tmp_path,
        [
            valid_row(**{"备注": None}),
            valid_row(**{"备注": ""}),
            valid_row(**{"备注": "   "}),
        ],
    )
    preview = excel.preview_import(path)
    assert preview.new_count == 1
    assert preview.duplicate_count == 2


def test_xls04_any_error_blocks_commit(tmp_path, excel):
    path = save_rows(
        tmp_path,
        [valid_row(), valid_row(**{"卖出日期": "2026/01/15"})],  # 第二行日期非法
    )
    preview = excel.preview_import(path)
    assert preview.has_errors
    assert preview.error_count == 1
    assert "无法识别" in preview.rows[1].reason
    assert preview.rows[1].row_number == 3  # Excel 实际行号（表头占第 1 行）
    with pytest.raises(ExcelError):
        excel.commit_import(preview)


def test_xls04_field_error_reasons(tmp_path, excel):
    path = save_rows(
        tmp_path,
        [
            valid_row(**{"卖出价格": "abc"}),
            valid_row(**{"卖出数量": 0}),
            valid_row(**{"卖出日期": date.today() + timedelta(days=1)}),
            valid_row(**{"转债代码": "12345"}),
        ],
    )
    preview = excel.preview_import(path)
    assert preview.error_count == 4
    assert "必须是数字" in preview.rows[0].reason
    assert "正整数" in preview.rows[1].reason
    assert "不能晚于今天" in preview.rows[2].reason
    assert "六位数字" in preview.rows[3].reason


def test_xls05_failed_write_rolls_back_all(tmp_path, excel, service):
    """XLS-05：单事务批量写入任一失败，全部回滚、无部分写入。"""
    good = SaleRecord("平银转债", "012345", date(2026, 1, 15), Decimal("130.00"), 10)
    bad = SaleRecord("违规记录", "999999", date(2026, 1, 15), Decimal("130.00"), 0)  # 绕过校验直接触发 CHECK
    with pytest.raises(RepositoryError):
        service.add_many([good, bad])
    assert service.list_records() == []


def test_xls06_shuffled_column_order(tmp_path, excel):
    """XLS-06：列顺序打乱但列名正确，正确识别。"""
    shuffled = ["卖出数量", "备注", "卖出价格", "转债代码", "转债名称", "卖出日期"]
    rows = [[10, "测试备注", "130.00", "012345", "平银转债", "2026-01-15"]]
    path = save_rows(tmp_path, rows, headers=shuffled)
    preview = excel.preview_import(path)
    assert preview.new_count == 1
    record = preview.rows[0].record
    assert record.bond_name == "平银转债"
    assert record.bond_code == "012345"
    assert record.sale_price == Decimal("130.00")
    assert record.quantity == 10
    assert record.note == "测试备注"


def test_xls08_lost_leading_zero_rejected(tmp_path, excel):
    """XLS-08：数字单元格丢失前导零 → 拒绝猜测补零。"""
    path = save_rows(tmp_path, [valid_row(**{"转债代码": 12345})])  # 012345 变数字
    preview = excel.preview_import(path)
    assert preview.error_count == 1
    assert "前导零" in preview.rows[0].reason
    assert "文本" in preview.rows[0].reason


def test_xls10_loss_batch_net_profit(tmp_path, excel):
    """XLS-10：含亏损的批次净收益正确计算（仅新增、含负收益、重复不计）。"""
    rows = [
        valid_row(),  # +300
        valid_row(**{"转债名称": "亏损转债", "转债代码": "111111", "卖出价格": "98.50"}),  # −15
        valid_row(),  # 重复，不计入
    ]
    path = save_rows(tmp_path, rows)
    preview = excel.preview_import(path)
    assert preview.net_profit == Decimal("285.00")


def test_blank_rows_skipped_with_real_row_numbers(tmp_path, excel):
    path = save_rows(
        tmp_path,
        [valid_row(), [None] * 6, [None] * 6, valid_row(**{"转债名称": "第二只"})],
    )
    preview = excel.preview_import(path)
    assert [row.row_number for row in preview.rows] == [2, 5]  # 空白行忽略且行号真实
    assert preview.new_count == 2


def test_missing_required_column_raises(tmp_path, excel):
    path = save_rows(
        tmp_path,
        [valid_row()],
        headers=["转债名称", "转债代码", "卖出日期", "卖出数量", "备注"],
    )
    with pytest.raises(ExcelError) as excinfo:
        excel.preview_import(path)
    assert "缺少必需列" in str(excinfo.value)
    assert "卖出价格" in str(excinfo.value)


def test_nonexistent_file_raises(tmp_path, excel):
    with pytest.raises(ExcelError):
        excel.preview_import(str(tmp_path / "no_such_file.xlsx"))


# ---- 提交（§4.3） ------------------------------------------------------------

def test_commit_success(tmp_path, excel, service):
    path = save_rows(
        tmp_path,
        [valid_row(), valid_row(**{"转债名称": "亏损转债", "转债代码": "111111", "卖出价格": "98.50"})],
    )
    preview = excel.preview_import(path)
    result = excel.commit_import(preview)

    assert result.added_count == 2
    assert result.skipped_count == 0
    assert result.net_profit == Decimal("285.00")
    assert len(service.list_records()) == 2


def test_commit_rechecks_duplicates_after_preview(tmp_path, excel, service):
    """§4.3：预览后数据库发生变化 → 提交时重新检查去重，以实际事务为准。"""
    path = save_rows(tmp_path, [valid_row()])
    preview = excel.preview_import(path)
    assert preview.new_count == 1

    service.create_record(SaleRecord("平银转债", "012345", date(2026, 1, 15), Decimal("130.00"), 10))
    result = excel.commit_import(preview)

    assert result.added_count == 0
    assert result.skipped_count == 1
    assert len(service.list_records()) == 1  # 手工那条，未重复写入
