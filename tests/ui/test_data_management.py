"""数据管理页测试（§4）：模板下载、导入全流程（monkeypatch 文件对话框）。"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

import bond_profit.ui.data_management as page_module
from bond_profit.application.excel_service import ExcelService
from bond_profit.application.sales_service import SalesService
from bond_profit.infrastructure.database import connect
from bond_profit.infrastructure.repository import SalesRepository
from bond_profit.ui.data_management import DataManagementPage

_HEADERS = ["转债名称", "转债代码", "卖出日期", "卖出价格", "卖出数量", "备注"]


@pytest.fixture()
def service(tmp_path):
    conn = connect(tmp_path / "app.db")
    yield SalesService(SalesRepository(conn))
    conn.close()


@pytest.fixture()
def page(service, qtbot):
    excel = ExcelService(service)
    widget = DataManagementPage(excel)
    qtbot.addWidget(widget)
    return widget


def save_valid_file(tmp_path, name="import.xlsx"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "卖出记录"
    sheet.append(_HEADERS)
    sheet.append(["平银转债", "012345", date(2026, 1, 15), "130.00", 10, ""])
    path = tmp_path / name
    workbook.save(path)
    return str(path)


# ---- 模板下载 -------------------------------------------------------------------

def test_download_template(page, tmp_path, monkeypatch, qtbot):
    target = tmp_path / "模板.xlsx"
    infos = []
    monkeypatch.setattr(
        page_module.QFileDialog, "getSaveFileName",
        lambda *a, **k: (str(target), ""),
    )
    monkeypatch.setattr(
        page_module.QMessageBox, "information",
        lambda parent, title, text: infos.append(title),
    )

    page.template_button.click()

    assert target.exists()
    workbook = load_workbook(target)
    assert "卖出记录" in workbook.sheetnames
    assert "填写说明" in workbook.sheetnames
    assert infos == ["模板已保存"]


def test_download_template_cancelled(page, tmp_path, monkeypatch):
    monkeypatch.setattr(
        page_module.QFileDialog, "getSaveFileName", lambda *a, **k: ("", "")
    )
    page.template_button.click()
    assert not (tmp_path / "模板.xlsx").exists()  # 未选路径时不产生文件


def test_download_template_write_failure(page, tmp_path, monkeypatch):
    warnings = []
    monkeypatch.setattr(
        page_module.QFileDialog, "getSaveFileName",
        lambda *a, **k: (str(tmp_path / "t.xlsx"), ""),
    )
    monkeypatch.setattr(
        page_module.QMessageBox, "warning",
        lambda parent, title, text: warnings.append(title),
    )

    def raise_oserror(_self, _path):
        raise OSError("disk full")

    monkeypatch.setattr(page_module.ExcelService, "generate_template", raise_oserror)
    page.template_button.click()

    assert warnings == ["模板保存失败"]


# ---- 导入 ---------------------------------------------------------------------

def test_import_success_emits_and_writes(page, service, tmp_path, monkeypatch, qtbot):
    path = save_valid_file(tmp_path)
    monkeypatch.setattr(
        page_module.QFileDialog, "getOpenFileName", lambda *a, **k: (path, "")
    )
    monkeypatch.setattr(
        page_module.QMessageBox, "information", lambda *a, **k: None
    )

    def fake_exec(self):
        # 走真实确认路径：commit_import → 信息框 → accept
        self._on_confirm()
        return True

    monkeypatch.setattr(page_module.ImportPreviewDialog, "exec", fake_exec)

    with qtbot.waitSignal(page.import_completed):
        page.import_button.click()

    assert len(service.list_records()) == 1


def test_import_file_level_error_shows_warning(page, service, tmp_path, monkeypatch):
    """缺必需列等文件级错误：提示中文原因，不打开预览、不写库。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "卖出记录"
    sheet.append(["转债名称", "转债代码", "卖出日期", "卖出数量", "备注"])
    sheet.append(["平银转债", "012345", "2026-01-15", 10, ""])
    path = tmp_path / "bad.xlsx"
    workbook.save(path)

    warnings = []
    monkeypatch.setattr(
        page_module.QFileDialog, "getOpenFileName", lambda *a, **k: (str(path), "")
    )
    monkeypatch.setattr(
        page_module.QMessageBox, "warning",
        lambda parent, title, text: warnings.append((title, text)),
    )

    page.import_button.click()

    assert warnings[0][0] == "无法导入"
    assert "缺少必需列" in warnings[0][1]
    assert len(service.list_records()) == 0


def test_import_cancelled_file_selection(page, service, monkeypatch):
    monkeypatch.setattr(
        page_module.QFileDialog, "getOpenFileName", lambda *a, **k: ("", "")
    )
    page.import_button.click()
    assert len(service.list_records()) == 0


def test_import_dialog_rejected_does_not_emit(page, service, tmp_path, monkeypatch, qtbot):
    """预览对话框取消 → 不发射 import_completed，不写库。"""
    path = save_valid_file(tmp_path)
    monkeypatch.setattr(
        page_module.QFileDialog, "getOpenFileName", lambda *a, **k: (path, "")
    )
    monkeypatch.setattr(
        page_module.ImportPreviewDialog, "exec", lambda self: False
    )

    page.import_button.click()

    assert len(service.list_records()) == 0
