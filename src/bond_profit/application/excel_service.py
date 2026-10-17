"""Excel 导入导出服务（文档 §4）。

模板生成（§4.1）、解析与逐行校验（§4.2）、标准化去重（§4.3）、
预览统计（§4.4）。解析为纯逻辑不写库；提交走单事务全有或全无，
且提交前重新检查去重（预览后数据库可能变化，最终以实际事务为准）。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook

from bond_profit.application.sales_service import SalesService
from bond_profit.domain.calculations import profit_per_record
from bond_profit.domain.money import yuan_to_cents
from bond_profit.domain.models import SaleRecord
from bond_profit.domain.validation import validate_record

_HEADERS = ("转债名称", "转债代码", "卖出日期", "卖出价格", "卖出数量", "备注")
_REQUIRED_HEADERS = _HEADERS[:5]  # 备注列可缺失
_SHEET_NAME = "卖出记录"
_CODE_PATTERN = re.compile(r"^[0-9]{6}$")
_INT_TEXT_PATTERN = re.compile(r"^[0-9]+$")


class ExcelError(Exception):
    """文件级错误（无法打开、缺失必需列等），信息可直接展示给用户。"""


@dataclass
class RowResult:
    """一行的解析结果（§4.4 逐行状态）。"""

    row_number: int  # Excel 实际行号（含表头偏移）
    status: str  # "new" / "duplicate" / "error"
    record: SaleRecord | None = None
    reason: str = ""  # 错误原因（仅 error 行）


@dataclass
class ImportPreview:
    rows: list[RowResult] = field(default_factory=list)
    new_count: int = 0
    duplicate_count: int = 0
    error_count: int = 0
    net_profit: Decimal = Decimal("0")  # 预计净收益变化：仅新增、含负收益（§4.4）

    @property
    def has_errors(self) -> bool:
        return self.error_count > 0

    @property
    def new_records(self) -> list[SaleRecord]:
        return [row.record for row in self.rows if row.status == "new"]


@dataclass
class CommitResult:
    """实际导入结果（§4.4）：以实际事务为准。"""

    added_count: int
    skipped_count: int  # 提交时重新检查发现重复而跳过
    net_profit: Decimal


class ExcelService:
    def __init__(self, sales_service: SalesService) -> None:
        self._sales_service = sales_service

    # ---- 模板（§4.1） ---------------------------------------------------------

    def generate_template(self, path: str) -> None:
        """生成官方模板：六列标准列名 + 填写说明页；数据区不放示例行。"""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = _SHEET_NAME
        sheet.append(list(_HEADERS))
        for row in range(2, 201):  # 预设常用格式：代码文本、日期/价格显示格式
            sheet.cell(row=row, column=2).number_format = "@"  # 代码保留前导零
            sheet.cell(row=row, column=3).number_format = "yyyy-mm-dd"
            sheet.cell(row=row, column=4).number_format = "0.00"
        for column, width in zip("ABCDEF", (24, 12, 14, 12, 12, 30)):
            sheet.column_dimensions[column].width = width

        guide = workbook.create_sheet("填写说明")
        guide.column_dimensions["A"].width = 100
        for line in (
            "填写说明",
            "1. 请在「卖出记录」表第 2 行起逐行填写，一行为一条记录；上方表头请勿修改。",
            "2. 转债名称：必填，不超过 50 个字符。",
            "3. 转债代码：必填，恰好六位数字。该列已设为文本格式，请勿改成数字格式，"
            "否则前导零会丢失（例如 012345 变成 12345），软件将拒绝导入。",
            "4. 卖出日期：必填，不能晚于今天。支持 Excel 日期单元格或 YYYY-MM-DD 文本。",
            "5. 卖出价格：必填，大于 0，最多两位小数。",
            "6. 卖出数量：必填，正整数。",
            "7. 备注：可选，最长 500 个字符；同一笔真实交易请用不同备注区分。",
            "8. 完全空白的行会被忽略；完全相同的两行只会导入第一行。",
        ):
            guide.append([line])
        workbook.save(path)

    # ---- 解析与预览（§4.2–§4.4） ----------------------------------------------

    def preview_import(self, path: str) -> ImportPreview:
        """解析整个文件并逐行校验、去重，返回预览。不做任何写入。"""
        rows = self._parse_rows(path)
        existing_keys = {
            self.record_key(record) for record in self._sales_service.list_records()
        }
        seen: set[tuple] = set()
        preview = ImportPreview(rows=rows)
        for row in rows:
            if row.status != "new":
                preview.error_count += 1
                continue
            key = self.record_key(row.record)
            if key in existing_keys or key in seen:
                row.status = "duplicate"  # 与库重复或文件内重复（§4.3）
                # record 保留供预览展示；new_records 只取 status=="new"，不会重复写入
                preview.duplicate_count += 1
                continue
            seen.add(key)
            preview.new_count += 1
            preview.net_profit += profit_per_record(
                row.record.sale_price, row.record.quantity
            )
        return preview

    def _parse_rows(self, path: str) -> list[RowResult]:
        header_map, data_rows = self._read_sheet(path)
        results: list[RowResult] = []
        for row_number, cells in data_rows:
            if all(
                value is None or (isinstance(value, str) and not value.strip())
                for value in cells
            ):
                continue  # 完全空白行跳过（§4.2）
            results.append(self._parse_row(row_number, cells, header_map))
        return results

    @staticmethod
    def _read_sheet(path: str) -> tuple[dict[str, int], list[tuple[int, tuple]]]:
        """读取数据表：优先「卖出记录」表，按列名匹配（列顺序不限）。

        返回 (列名→列号, [(Excel 行号, 单元格元组)])。
        """
        try:
            workbook = load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:
            raise ExcelError(f"无法打开 Excel 文件：{exc}") from exc
        try:
            sheet = next(
                (s for s in workbook.worksheets if s.title == _SHEET_NAME),
                workbook.worksheets[0],
            )
            all_rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()

        if not all_rows:
            raise ExcelError("表格为空，缺少表头行")
        header_map: dict[str, int] = {}
        for index, value in enumerate(all_rows[0]):
            if isinstance(value, str) and value.strip():
                header_map.setdefault(value.strip(), index)
        missing = [name for name in _REQUIRED_HEADERS if name not in header_map]
        if missing:
            raise ExcelError("缺少必需列：" + "、".join(missing))
        data_rows = [
            (number, row) for number, row in enumerate(all_rows[1:], start=2)
        ]
        return header_map, data_rows

    def _parse_row(
        self, row_number: int, cells: tuple, header_map: dict[str, int]
    ) -> RowResult:
        """把一行单元格转换为标准记录；单元格级错误给出行号+中文原因。"""
        errors: list[str] = []

        def raw(name: str):
            index = header_map.get(name)
            return cells[index] if index is not None and index < len(cells) else None

        # 名称（§4.2：文本；名称去首尾空格在标准化时进行）
        name_value = raw("转债名称")
        bond_name = ""
        if name_value is None or (isinstance(name_value, str) and not name_value.strip()):
            pass  # 留空，由最终校验报「名称不能为空」
        elif isinstance(name_value, str):
            bond_name = name_value.strip()
        else:
            errors.append("转债名称必须是文本")

        # 代码：数值单元格 = Excel 已丢失前导零，不猜测补零（§4.2 / XLS-08）
        code_value = raw("转债代码")
        bond_code = ""
        if code_value is None or (isinstance(code_value, str) and not code_value.strip()):
            pass  # 留空，由最终校验报错
        elif isinstance(code_value, bool):
            errors.append("转债代码必须是六位数字文本")
        elif isinstance(code_value, (int, float)):
            errors.append(
                "转债代码列是数字格式，前导零已丢失，请修正源文件为文本格式后重新导入"
            )
        elif isinstance(code_value, str):
            bond_code = code_value.strip()
            if not _CODE_PATTERN.fullmatch(bond_code):
                bond_code = ""
                errors.append("转债代码必须是恰好六位数字")
        else:
            errors.append("转债代码必须是六位数字文本")

        # 日期：Excel 日期单元格或 YYYY-MM-DD 文本（XLS-07）
        date_value = raw("卖出日期")
        sale_date: date | None = None
        if date_value is None or (isinstance(date_value, str) and not date_value.strip()):
            errors.append("卖出日期不能为空")
        elif isinstance(date_value, datetime):
            sale_date = date_value.date()
        elif isinstance(date_value, date):
            sale_date = date_value
        elif isinstance(date_value, str):
            try:
                sale_date = date.fromisoformat(date_value.strip())
            except ValueError:
                errors.append("卖出日期无法识别，请使用 YYYY-MM-DD 文本或 Excel 日期")
        else:
            errors.append("卖出日期无法识别，请使用 YYYY-MM-DD 文本或 Excel 日期")

        # 价格：有效数值或可精确解析的十进制文本（>0、两位小数由最终校验把关）
        price_value = raw("卖出价格")
        sale_price: Decimal | None = None
        if price_value is None or (isinstance(price_value, str) and not price_value.strip()):
            errors.append("卖出价格不能为空")
        elif isinstance(price_value, bool):
            errors.append("卖出价格必须是数字")
        elif isinstance(price_value, (int, float)):
            sale_price = Decimal(str(price_value))
        elif isinstance(price_value, str):
            try:
                sale_price = Decimal(price_value.strip())
            except InvalidOperation:
                errors.append("卖出价格必须是数字")
        else:
            errors.append("卖出价格必须是数字")

        # 数量：正整数（XLS-04）
        quantity_value = raw("卖出数量")
        quantity: int | None = None
        if quantity_value is None or (
            isinstance(quantity_value, str) and not quantity_value.strip()
        ):
            errors.append("卖出数量不能为空")
        elif isinstance(quantity_value, bool):
            errors.append("卖出数量必须是正整数")
        elif isinstance(quantity_value, int):
            quantity = quantity_value
        elif isinstance(quantity_value, float) and quantity_value.is_integer():
            quantity = int(quantity_value)
        elif isinstance(quantity_value, str) and _INT_TEXT_PATTERN.fullmatch(
            quantity_value.strip()
        ):
            quantity = int(quantity_value)
        else:
            errors.append("卖出数量必须是正整数")

        # 备注：空单元格/空串/纯空格统一为空（§4.3）
        note_value = raw("备注")
        note = ""
        if note_value is None:
            note = ""
        elif isinstance(note_value, str):
            note = note_value.strip()
        else:
            errors.append("备注必须是文本")

        if errors or sale_date is None or sale_price is None or quantity is None:
            if not errors:  # 防御：缺字段而未给原因
                errors.append("行数据不完整")
            return RowResult(
                row_number=row_number, status="error", reason="；".join(errors)
            )

        record = SaleRecord(
            bond_name=bond_name,
            bond_code=bond_code,
            sale_date=sale_date,
            sale_price=sale_price,
            quantity=quantity,
            note=note,
        )
        field_errors = validate_record(record)
        if field_errors:
            return RowResult(
                row_number=row_number, status="error", reason="；".join(field_errors)
            )
        return RowResult(row_number=row_number, status="new", record=record)

    # ---- 提交（§4.3 事务） ------------------------------------------------------

    def commit_import(self, preview: ImportPreview) -> CommitResult:
        """提交预览中的新增记录：重新检查去重后单事务写入，全有或全无。"""
        if preview.has_errors:  # 存在任一错误时整批不可提交（§4.3）
            raise ExcelError("存在错误行，请修正后重新导入")
        existing_keys = {
            self.record_key(record) for record in self._sales_service.list_records()
        }
        to_add: list[SaleRecord] = []
        skipped = 0
        for record in preview.new_records:
            key = self.record_key(record)
            if key in existing_keys:  # 预览后库中已出现相同记录（§4.3 重新检查）
                skipped += 1
                continue
            existing_keys.add(key)
            to_add.append(record)
        added = self._sales_service.add_many(to_add)  # 失败抛 RepositoryError，全部未写
        net_profit = sum(
            (profit_per_record(r.sale_price, r.quantity) for r in added),
            Decimal("0"),
        )
        return CommitResult(
            added_count=len(added), skipped_count=skipped, net_profit=net_profit
        )

    @staticmethod
    def record_key(record: SaleRecord) -> tuple:
        """导入专用去重键（§4.3）：标准化后的六字段。"""
        return (
            record.bond_name.strip(),
            record.bond_code.strip(),
            record.sale_date.isoformat(),
            yuan_to_cents(record.sale_price),
            int(record.quantity),
            record.note.strip(),
        )
